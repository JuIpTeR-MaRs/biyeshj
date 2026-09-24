from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING

OUT=r'D:\biyesheji\基于区块链的家庭监护消费管理系统_开题报告.docx'
doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5); sec.left_margin=Cm(2.8); sec.right_margin=Cm(2.6)

styles=doc.styles
styles['Normal'].font.name='宋体'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); styles['Normal'].font.size=Pt(12)
styles['Normal'].paragraph_format.line_spacing=1.5
for sname,size,bold in [('Title',18,True),('Heading 1',14,True),('Heading 2',12,True)]:
    s=styles[sname]; s.font.name='黑体'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); s.font.size=Pt(size); s.font.bold=bold; s.font.color.rgb=None
    s.paragraph_format.space_before=Pt(8); s.paragraph_format.space_after=Pt(5)

def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_text(cell,text,bold=False,align=WD_ALIGN_PARAGRAPH.LEFT,size=11):
    cell.text=''; p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.line_spacing=1.3
    r=p.add_run(text); r.bold=bold; r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(size)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_borders(table,color='B7B7B7',sz='6'):
    tblPr=table._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),sz); e.set(qn('w:color'),color); borders.append(e)
    tblPr.append(borders)

def body(text, indent=True):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.5; p.paragraph_format.space_after=Pt(2)
    if indent: p.paragraph_format.first_line_indent=Cm(0.74)
    r=p.add_run(text); r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)
    return p

def item(num,title,text):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.5; p.paragraph_format.space_after=Pt(2); p.paragraph_format.first_line_indent=Cm(0.74)
    r=p.add_run(f'{num}. {title}'); r.bold=True; r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)
    r=p.add_run(text); r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(6)
r=p.add_run('广州商学院本科毕业论文（设计）开题报告'); r.bold=True; r.font.name='黑体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); r.font.size=Pt(18)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14)
r=p.add_run('（学生用表）'); r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(12)

info=doc.add_table(rows=4, cols=4); info.alignment=WD_TABLE_ALIGNMENT.CENTER; info.autofit=False
widths=[Cm(2.5),Cm(6.0),Cm(2.5),Cm(6.0)]
labels=[('学院','现代信息产业学院','专业','软件工程'),('学生姓名','________________','学号','________________'),('班级','________________','指导教师','________________'),('论文（设计）题目','基于区块链的家庭监护消费管理系统的设计与实现','','')]
for i,row in enumerate(info.rows):
    for j,c in enumerate(row.cells): c.width=widths[j]
    vals=labels[i]
    if i==3:
        row.cells[1].merge(row.cells[3]); vals=('论文（设计）题目','基于区块链的家庭监护消费管理系统的设计与实现','','')
    set_cell_text(row.cells[0],vals[0],True,WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(row.cells[0],'E7E6E6')
    set_cell_text(row.cells[1],vals[1],False,WD_ALIGN_PARAGRAPH.CENTER)
    if i<3:
        set_cell_text(row.cells[2],vals[2],True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_shading(row.cells[2],'E7E6E6')
        set_cell_text(row.cells[3],vals[3],False,WD_ALIGN_PARAGRAPH.CENTER)
set_borders(info)

doc.add_paragraph()
doc.add_heading('一 本题目的意义及国内外研究状况',level=1)
body('随着移动支付、网络购物和数字服务不断进入家庭生活，未成年人和部分需要数字协助的老年人能够更加便捷地完成购物、充值及生活缴费，但也面临消费判断能力不足、超额消费发现不及时、家庭成员审批不便以及关键处理记录难以核验等问题。现有个人记账和账单统计工具主要服务于事后分析，难以形成消费发生前的额度判断、受限类别识别和家庭协同审批流程。因此，设计一套兼顾规则控制、监护审批和可信追溯的家庭消费管理系统，具有明确的实际应用价值。')
body('在家庭数字消费研究方面，蒋媛媛等分析了社区居家养老服务对老年家庭数字消费的影响，指出数字能力、移动支付使用和家庭支持条件会影响老年人的数字消费行为[1]。冀嘉豪、吴华安从代际数字反哺角度讨论老年群体消费升级，为家庭成员参与数字协助提供了参考[2]；陈婷等研究农村老年人的数字支付抵制行为，说明易用性、风险认知和使用心理是老年支付服务需要考虑的重要因素[3]。未成年人方面，季为民、孙芳分析了未成年人互联网运用状况[4]，陈玉真研究其互联网风险认知和行为表现[5]，潘俊则从法律效力角度讨论未成年人网络消费及法定代理人的同意与追认问题[6]。这些研究表明，老年人和未成年人的管理方式需要有所区别：老年人场景应尊重本人意愿并以协助授权为基础，未成年人场景应结合监护关系、额度限制和审批规则进行管理。')
body('在区块链和智能合约研究方面，陈晶等对区块链扩展技术及性能问题进行了总结[7]；巩坪等提出基于智能合约的权限传递访问控制模型，为授权关系的建立、传递和验证提供了技术思路[8]；刘炜等研究基于区块链的多授权机构访问控制方法，说明区块链能够支持多主体权限策略的协同管理[9]。韩永征、胡春强提出动态信任委托访问控制模型，其委托授权、撤销和审计追踪机制可为家庭成员授权设计提供辅助参考[10]。但该研究面向体域网，且底层共识描述存在衔接不够清楚的问题，本文只借鉴其授权与审计思想，不直接采用其底层架构。')
body('在数据存储与可信追溯方面，刘春等提出链上链下结合的数据安全存储与访问控制方案，将完整数据存储在链下，将哈希、签名和关键操作记录写入区块链[11]；潘恒等总结了典型区块链存储与查询技术，为控制链上存储规模和提高查询效率提供了依据[12]；宋靖文等的身份隐私保护研究说明链上数据设计应减少身份信息暴露[13]。付豪完成了基于区块链的可信交易系统设计与实现，为本课题的需求分析、模块划分、交易状态管理、统计报表和系统测试提供了工程参考[14]。钱鹏等对智能合约安全漏洞检测技术进行了综述，可用于指导合约权限、状态转换和重入等风险的测试[15]。')
body('综合现有研究，家庭数字消费、未成年人保护、老年数字协助、区块链访问控制和链上存证均已形成一定研究基础，但将消费额度、商户类别、账户状态、家庭授权和审批结果整合为完整业务流程的研究相对较少。本课题拟形成“消费请求—规则判断—自动通过或监护审批—模拟执行—链上存证—记录核验”的闭环，以React构建前端界面，Node.js实现后端服务，MySQL保存业务数据，Solidity智能合约保存关键状态和数据摘要，验证区块链在家庭消费辅助管理场景中的应用方式。')
body('本课题的意义主要体现在两个方面。应用方面，通过事前规则判断和必要的监护审批，帮助家庭及时发现并处理超额或受限消费，提高消费管理的透明度；通过保存审批原因、处理结果和链上凭证，增强关键记录的可追溯性。工程方面，研究前端、后端、数据库与智能合约之间的协同方式，重点处理权限检查、交易状态转换、链上链下数据一致性和异常恢复，为智能合约在家庭消费辅助管理场景中的应用提供可运行、可测试的实践参考。')

doc.add_heading('二 研究内容',level=1)
body('本研究面向未成年人和需要消费协助的老年人，设计并实现家庭监护消费管理系统。系统采用模拟消费数据，不接入真实资金扣划，主要研究内容如下：')
item(1,'用户需求与总体架构设计。','分析消费成员、监护或协助成员及平台管理员的功能需求，明确角色权限和数据访问边界。采用前后端分离架构，划分账户管理、家庭关联、消费规则、交易审批、链上存证和统计查询等模块，完成业务流程、数据库、接口和合约数据结构设计。')
item(2,'家庭关联与差异化权限管理。','实现家庭成员邀请、关联确认、权限配置、关系变更及历史查询。老人场景以本人确认的协助授权为基础，支持授权范围、有效期和撤销；未成年人场景以监护账户关联为基础，支持消费规则配置与审批。系统分别检查家庭归属与操作权限，防止跨家庭访问和越权审批。')
item(3,'消费规则判断与额度控制。','根据单次消费金额、周期累计额度、商户类别和账户状态建立规则判断机制。账户正常且满足规则的消费自动通过；超过自动放行阈值或属于需确认类别的消费进入审批；账户冻结、关联失效等请求停止执行。记录规则版本和触发原因，并通过数据库事务处理并发额度占用与释放。')
item(4,'监护审批与交易状态管理。','实现待审批交易查询、详情查看、审批通过、拒绝、取消和超时处理。设计待判定、待审批、已批准、已拒绝、已执行和已失效等状态，限定合法状态转换。执行前重新检查账户、关联和规则状态，避免失效授权、旧规则或重复请求继续执行。')
item(5,'智能合约存证与链上链下一致性处理。','使用Solidity实现关键交易状态、审批结果、操作主体、时间戳和数据摘要的记录与查询，限制合约写入权限及状态转换顺序。MySQL保存消费明细和统计所需数据，区块链保存核验凭证。通过交易回执查询、幂等控制、失败重试和定期对账处理数据库状态与链上确认状态的差异。')
item(6,'消费查询、统计分析与移动端适配。','实现按时间、家庭成员、商户类别和交易状态查询记录，展示消费总额、类别分布、额度使用和待审批数量。采用响应式布局适配手机和桌面浏览器，优化消费提交、审批确认、异常提示和结果查询等核心操作。')
item(7,'系统测试与结果分析。','围绕自动放行、超额审批、受限类别、账户冻结、授权撤销、重复提交、并发额度更新和上链失败等场景开展功能与安全测试；记录业务接口响应时间、处理成功率和链上确认时延，分析系统正确性、稳定性和使用体验。')

doc.add_heading('三 研究方法 手段及步骤',level=1)
item(1,'文献研究法。','检索并阅读家庭数字消费、未成年人网络消费、老年数字支付、区块链访问控制、链上链下存储和智能合约安全等文献，归纳现有研究的主要方法、适用范围与不足，为需求分析和技术选型提供依据。')
item(2,'需求分析法。','结合家庭消费场景开展同类应用调研和场景走查，使用用例图、业务流程图及权限矩阵明确角色职责、功能边界和异常情况，区分老人协助授权与未成年人监护管理的规则。')
item(3,'系统建模法。','使用实体关系图设计用户、家庭关系、消费规则、交易、审批及链上凭证等数据实体；使用状态图描述消费请求的生命周期；明确前端、后端、MySQL和智能合约之间的数据流转关系。')
item(4,'原型设计与迭代开发法。','先设计消费申请、规则设置和监护审批等页面原型，再依次实现账户与家庭关系、消费规则与审批、智能合约存证、统计分析及移动端适配。后端使用Node.js与Express提供接口，通过ethers.js与本地EVM网络交互。')
item(5,'实验测试法。','构造不同金额、商户类别、账户状态和授权条件的模拟交易，验证规则判断和审批结果；开展跨家庭访问、越权审批、重复提交、并发更新和链上失败恢复测试；记录测试环境、数据规模、接口时延和链上确认时延。')
item(6,'反馈优化法。','依据功能测试、场景走查和使用反馈修正规则、状态转换和异常处理逻辑，改进页面提示与移动端操作，确保系统实际行为、测试结果和论文描述一致。')
body('实施步骤为：完成文献调研和需求分析；完成系统架构、数据库、接口与合约设计；实现家庭关联和消费规则；实现审批流程和状态管理；完成智能合约部署及链上链下联调；实现统计分析和移动端适配；开展测试、优化并完成论文。')

doc.add_heading('四 进度安排',level=1)
rows=[('第1—2周','完成文献调研、选题分析和需求梳理，撰写开题报告。'),('第3—4周','完成页面原型、系统架构、数据库、接口和智能合约设计，搭建开发环境。'),('第5—6周','实现账户管理、家庭关联、授权管理和消费规则配置。'),('第7—8周','实现消费判断、自动放行、监护审批和模拟执行，完成核心业务流程。'),('第9—10周','完成智能合约集成、存证查询、失败重试、统计分析和移动端适配。'),('第11—12周','开展功能、安全、并发和性能测试，修复问题并整理实验数据。'),('第13—14周','完成论文初稿、部署说明和用户操作文档，根据指导意见修改。'),('第15—16周','完成系统与论文定稿，准备演示数据和答辩材料。')]
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
set_cell_text(t.rows[0].cells[0],'阶段',True,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(t.rows[0].cells[1],'主要任务',True,WD_ALIGN_PARAGRAPH.CENTER)
for c in t.rows[0].cells: set_cell_shading(c,'D9EAF7')
for a,b in rows:
    cells=t.add_row().cells; set_cell_text(cells[0],a,False,WD_ALIGN_PARAGRAPH.CENTER); set_cell_text(cells[1],b)
set_borders(t)

doc.add_heading('五 主要参考文献',level=1)
refs=[
'[1] 蒋媛媛,温悦,陈瑞.社区居家养老与老年家庭数字消费[J].贵州财经大学学报,2025(6):63-70.',
'[2] 冀嘉豪,吴华安.代际数字反哺对老年群体消费升级的影响[J].攀枝花学院学报,2026,43(4):23-35.',
'[3] 陈婷,吴江,沈校亮,等.农村老年人数字支付抵制行为的心理归因及其作用机制研究[J].管理学报,2024,21(1):107-116,126.',
'[4] 季为民,孙芳.未成年人互联网运用状况问题研究[J].青年发展论坛,2023,33(4):25-33.',
'[5] 陈玉真.未成年人的互联网风险认知与行为呈现[J].北京青年研究,2023,32(3):60-71.',
'[6] 潘俊.未成年人网络消费行为的法律效力[J].太原理工大学学报(社会科学版),2023,41(5):59-66.',
'[7] 陈晶,杨浩,何琨,等.区块链扩展技术现状与展望[J].软件学报,2024,35(2):828-851.',
'[8] 巩坪,王九如,宋万水,等.基于智能合约的物联网权限传递访问控制模型[J].郑州大学学报(理学版),2023,55(3):28-33.',
'[9] 刘炜,李淑培,田钊,等.基于区块链的去中心化多授权机构访问控制方法[J].郑州大学学报(理学版),2025,57(5):46-53.',
'[10] 韩永征,胡春强.基于区块链的体域网动态信任委托访问控制研究[J].重庆大学学报,2026,49(6):93-102.',
'[11] 刘春,张凌浩,蒋昌松,等.基于联盟链的电力数据安全存储与访问控制方案[J].电子科技大学学报,2026,55(3):455-463.',
'[12] 潘恒,钱海洋,姚中原,等.典型区块链存储与查询技术综述[J].郑州大学学报(理学版),2022,54(6):34-50.',
'[13] 宋靖文,张大伟,韩旭,等.区块链中可监管的身份隐私保护方案[J].软件学报,2023,34(7):3292-3312.',
'[14] 付豪.基于区块链的虚拟电厂可信交易系统设计与实现[D].北京:北京邮电大学,2025.',
'[15] 钱鹏,刘振广,何钦铭,等.智能合约安全漏洞检测技术研究综述[J].软件学报,2022,33(8):3059-3085.'
]
for ref in refs:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(0.74); p.paragraph_format.first_line_indent=Cm(-0.74); p.paragraph_format.line_spacing=1.35; p.paragraph_format.space_after=Pt(1)
    r=p.add_run(ref); r.font.name='宋体'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(10.5)

doc.add_heading('六 指导教师意见',level=1)
t=doc.add_table(rows=2,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
set_cell_text(t.rows[0].cells[0],'（由指导教师填写）\n\n\n\n\n',False)
set_cell_text(t.rows[1].cells[0],'指导教师签名：________________        日期：______年____月____日',False,WD_ALIGN_PARAGRAPH.RIGHT)
set_borders(t)

doc.add_paragraph()
doc.add_heading('七 教研室或学院审核意见',level=1)
t=doc.add_table(rows=2,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
set_cell_text(t.rows[0].cells[0],'（由教研室或学院填写）\n\n\n\n',False)
set_cell_text(t.rows[1].cells[0],'负责人签名：________________          日期：______年____月____日',False,WD_ALIGN_PARAGRAPH.RIGHT)
set_borders(t)

# footer page number field
for section in doc.sections:
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

doc.core_properties.title='基于区块链的家庭监护消费管理系统的设计与实现 开题报告'
doc.core_properties.subject='广州商学院本科毕业论文（设计）开题报告'
doc.save(OUT)
print(OUT)
